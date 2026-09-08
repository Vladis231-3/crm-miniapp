"""XLSX formula injection: свободный текст (заметки, названия расходов)
попадает в xlsx без нейтрализации — Excel выполнит =,+,-,@ как формулу.

Фикс: app/exports._append_sheet нейтрализует все строковые ячейки.
"""

from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


class XlsxNeutralizationTests(unittest.TestCase):
    def test_helper_prefixes_dangerous_leads(self) -> None:
        from app.exports import _neutralize_xlsx_value

        for dangerous in (
            "=1+1",
            "=HYPERLINK(\"http://evil\",\"x\")",
            "+cmd",
            "-2+3",
            "@SUM(1+1)",
            "\t=1",
            "\r=1",
        ):
            with self.subTest(value=dangerous):
                self.assertEqual(
                    _neutralize_xlsx_value(dangerous), "'" + dangerous
                )

    def test_helper_leaves_safe_values_alone(self) -> None:
        from app.exports import _neutralize_xlsx_value

        self.assertEqual(_neutralize_xlsx_value("Мойка"), "Мойка")
        self.assertEqual(_neutralize_xlsx_value(""), "")
        self.assertEqual(_neutralize_xlsx_value("  =не сначала"), "  =не сначала")
        self.assertEqual(_neutralize_xlsx_value(123), 123)
        self.assertIsNone(_neutralize_xlsx_value(None))

    def test_append_sheet_roundtrip_neutralizes(self) -> None:
        from openpyxl import Workbook, load_workbook

        from app.exports import _append_sheet

        workbook = Workbook()
        workbook.remove(workbook.active)
        _append_sheet(
            workbook,
            "Расходы",
            ["Дата", "Название"],
            [["01.09.2026", "=HYPERLINK(\"http://evil\",\"жми\")"], ["02.09.2026", "Химия"]],
        )
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        sheet = load_workbook(buffer)["Расходы"]
        values = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(values[0], ["Дата", "Название"])
        self.assertEqual(values[1][1], "'=HYPERLINK(\"http://evil\",\"жми\")")
        self.assertEqual(values[2][1], "Химия")


if __name__ == "__main__":
    unittest.main()
